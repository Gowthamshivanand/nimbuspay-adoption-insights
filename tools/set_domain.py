"""One command before the pitch: put Gowtham's real link and QR code into the pitch files.

Usage (from the pack folder, the repo folder, or anywhere with --pack):
    python tools/set_domain.py https://your-domain.example
    python tools/set_domain.py https://your-domain.example --repo https://github.com/you/nimbuspay-adoption-insights
    python tools/set_domain.py https://your-domain.example --pack /path/to/GOWTHAM_TPM_PACK --no-pdf

What it writes (new files, originals untouched; every output ends in _final):
    extras/qr_final.png, extras/qr_final.svg          QR for the live link (and qr_repo_final.* with --repo)
    deck/NimbusPay_TPM_Adoption_Mirage_GowthamS_final.pptx (+ .pdf if LibreOffice is installed)
        slide 10: QR image, QR label and the "Try it" link box
    extras/LEAVE_BEHIND_TPM_final.pdf                  QR + links box (needs Chrome, Edge or Chromium)
    extras/CHEAT_SHEET_TPM_final.pdf                   links box on page 3

Needs: pip install segno python-pptx. Nothing is uploaded or published; it only writes local files.
"""
import argparse, glob, os, re, shutil, subprocess, sys, tempfile, urllib.parse

DECK = "deck/NimbusPay_TPM_Adoption_Mirage_GowthamS.pptx"
TEMPLATES = {"leave": "extras/_build/leave_tpm.html", "cheat": "extras/_build/cheat_tpm.html"}
OUTS = {"leave": "extras/LEAVE_BEHIND_TPM_final.pdf", "cheat": "extras/CHEAT_SHEET_TPM_final.pdf"}


def find_pack(explicit):
    if explicit:
        return os.path.abspath(explicit)
    here = os.path.dirname(os.path.abspath(__file__))
    for start in (os.getcwd(), here):
        d = start
        for _ in range(5):
            if os.path.exists(os.path.join(d, DECK)):
                return d
            d = os.path.dirname(d)
    sys.exit(f"Could not find {DECK}. Run from the pack folder or pass --pack PATH.")


def check_url(u, what):
    p = urllib.parse.urlparse(u)
    if p.scheme not in ("http", "https") or not p.netloc or "." not in p.netloc:
        sys.exit(f"{what} must be a full link like https://your-domain.example (got {u!r})")
    if re.search(r"YOUR-|example\.invalid|<|>", u, re.I):
        sys.exit(f"{what} still looks like a placeholder: {u!r}")
    return u.rstrip("/")


def show(u):
    return re.sub(r"^https?://", "", u)


def make_qr(url, base):
    import segno
    q = segno.make(url, error="m")
    q.save(base + ".png", scale=10, border=4, dark="#2B2B2B")
    q.save(base + ".svg", scale=10, border=4, dark="#2B2B2B", title=url)
    return base + ".png", base + ".svg"


def inline_svg(svg_path):
    s = re.sub(r"<\?xml[^>]*\?>", "", open(svg_path, encoding="utf-8").read())
    s = re.sub(r'width="\d+" height="\d+"', 'width="100%" height="100%"', s, count=1)
    if "viewBox" not in s:
        m = re.search(r'<svg[^>]*>', s)
        s = s.replace(m.group(0), m.group(0)[:-1] + ' viewBox="0 0 330 330">', 1) if m else s
    return s


def find_browser():
    for name in ("google-chrome", "google-chrome-stable", "chromium", "chromium-browser", "chrome", "msedge", "microsoft-edge"):
        p = shutil.which(name)
        if p:
            return p
    for p in ("/Applications/Google Chrome.app/Contents/MacOS/Google Chrome", "/Applications/Microsoft Edge.app/Contents/MacOS/Microsoft Edge",
              r"C:\Program Files\Google\Chrome\Application\chrome.exe", r"C:\Program Files (x86)\Google\Chrome\Application\chrome.exe",
              r"C:\Program Files (x86)\Microsoft\Edge\Application\msedge.exe", r"C:\Program Files\Microsoft\Edge\Application\msedge.exe"):
        if os.path.exists(p):
            return p
    return None


def html_to_pdf(html_path, pdf_path):
    b = find_browser()
    if b:
        r = subprocess.run([b, "--headless", "--no-sandbox", "--disable-gpu", "--no-pdf-header-footer",
                            f"--print-to-pdf={pdf_path}", "file:///" + os.path.abspath(html_path).replace(os.sep, "/").lstrip("/")],
                           capture_output=True, timeout=120)
        if r.returncode == 0 and os.path.exists(pdf_path):
            return True
    try:
        from playwright.sync_api import sync_playwright
        with sync_playwright() as p:
            br = p.chromium.launch(); pg = br.new_page(); pg.goto("file:///" + os.path.abspath(html_path).lstrip("/"))
            pg.pdf(path=pdf_path, prefer_css_page_size=True, print_background=True); br.close()
        return True
    except Exception:
        return False


def fill_block(html, marker, content):
    pat = re.compile(rf"<!--{marker}-->.*?<!--/{marker}-->", re.S)
    if not pat.search(html):
        raise ValueError(f"marker {marker} not found")
    return pat.sub(lambda m: f"<!--{marker}-->{content}<!--/{marker}-->", html)


def do_html(kind, pack, url, repo, qr_svg):
    src = os.path.join(pack, TEMPLATES[kind])
    if not os.path.exists(src):
        print(f"  skipped {kind}: template {TEMPLATES[kind]} not found"); return None
    h = open(src, encoding="utf-8").read()
    code = show(repo) if repo else None
    if kind == "leave":
        links = (f"<b>Try it live and read the code</b><br>Live demo: <code>{show(url)}</code><br>"
                 + (f"Code: <code>{code}</code><br>" if code else "") + "<span style='font-size:7pt;color:#555'>Scan the QR code or type the link.</span>")
        qr = inline_svg(qr_svg)
    else:
        links = f"Live demo: <b>{url}</b><br>" + (f"Code: <b>{code}</b><br>" if code else "") + "<span style='color:#555'>Same QR as slide 10 and the leave-behind.</span>"
        qr = inline_svg(qr_svg)
    h = fill_block(fill_block(h, "QR", qr), "LINKS", links)
    with tempfile.TemporaryDirectory() as td:
        tmp = os.path.join(td, f"{kind}_final.html"); open(tmp, "w", encoding="utf-8").write(h)
        out = os.path.join(pack, OUTS[kind])
        if html_to_pdf(tmp, out):
            return out
        keep = out[:-4] + ".html"; shutil.copy(tmp, keep)
        print(f"  no Chrome/Edge/Chromium found: wrote {keep}; open it in a browser and Print > Save as PDF")
        return keep


def do_deck(pack, url, repo, qr_png, make_pdf):
    from pptx import Presentation
    src = os.path.join(pack, DECK); out = src[:-5] + "_final.pptx"
    prs = Presentation(src); s = prs.slides[9]
    shapes = {sh.name: sh for sh in s.shapes}
    for need in ("qr-image", "qr-label", "demo-url"):
        if need not in shapes:
            sys.exit(f"Slide 10 has no shape named {need!r}; was the deck layout changed?")
    pic = shapes["qr-image"]; rid = pic._element.blipFill.blip.rEmbed
    pic.part.related_part(rid)._blob = open(qr_png, "rb").read()
    def set_para(p, text, link=None):
        runs = p.runs
        runs[0].text = text
        for r in runs[1:]: r.text = ""
        if link: runs[0].hyperlink.address = link
    set_para(shapes["qr-label"].text_frame.paragraphs[0], "Scan to try it")
    ps = shapes["demo-url"].text_frame.paragraphs
    set_para(ps[0], "Try it live")
    set_para(ps[1], show(url), url)
    if len(ps) > 2:
        set_para(ps[2], f"Code: {show(repo)}" if repo else "Code: linked from the site", repo)
        if repo and len(show(repo)) > 38:
            from pptx.util import Pt
            ps[2].runs[0].font.size = Pt(11)
    prs.save(out)
    pdf = None
    so = shutil.which("soffice") or shutil.which("libreoffice")
    if make_pdf and so:
        r = subprocess.run([so, "--headless", "--convert-to", "pdf", "--outdir", os.path.dirname(out), out], capture_output=True, timeout=300)
        pdf = out[:-5] + ".pdf" if r.returncode == 0 and os.path.exists(out[:-5] + ".pdf") else None
    elif make_pdf:
        print("  LibreOffice not found: open the _final.pptx in PowerPoint and export a PDF if you need one")
    return out, pdf


def main():
    a = argparse.ArgumentParser(description="Write Gowtham's real link and QR code into slide 10, the leave-behind and the cheat sheet (_final files).")
    a.add_argument("url", help="live demo link, e.g. https://gowtham.dev/nimbuspay")
    a.add_argument("--repo", help="public code link, e.g. https://github.com/<you>/nimbuspay-adoption-insights")
    a.add_argument("--pack", help="folder that contains deck/ and extras/ (auto-detected)")
    a.add_argument("--no-pdf", action="store_true", help="skip the deck PDF export")
    x = a.parse_args()
    url = check_url(x.url, "URL"); repo = check_url(x.repo, "--repo") if x.repo else None
    pack = find_pack(x.pack); ex = os.path.join(pack, "extras"); os.makedirs(ex, exist_ok=True)
    print(f"Pack: {pack}\nLink: {url}" + (f"\nCode: {repo}" if repo else ""))
    png, svg = make_qr(url, os.path.join(ex, "qr_final")); written = [png, svg]
    if repo:
        written += list(make_qr(repo, os.path.join(ex, "qr_repo_final")))
    deck, pdf = do_deck(pack, url, repo, png, not x.no_pdf); written += [deck] + ([pdf] if pdf else [])
    for k in ("leave", "cheat"):
        r = do_html(k, pack, url, repo, svg)
        if r: written.append(r)
    print("Wrote:"); [print("  " + os.path.relpath(w, pack)) for w in written]
    print("Check slide 10, the leave-behind and page 3 of the cheat sheet, then scan the QR with your phone.")


if __name__ == "__main__":
    main()
